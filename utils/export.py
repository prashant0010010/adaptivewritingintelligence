"""Report building and export (Markdown, TXT, DOCX). Observations, estimates and preferences are labelled separately."""
from __future__ import annotations

import io
import json
import re
from datetime import datetime

from core.profile_learning import DIMENSIONS
from core.writing_analysis import build_genre_profile
from database.repository import Repository
from utils.text_processing import change_summary


def build_report_markdown(repo: Repository, profile_id: int, max_revisions: int = 5) -> str:
    prof = repo.get_profile(profile_id) or {"name": "Unknown"}
    counts = repo.counts(profile_id)
    L = [f"# Writing report: {prof['name']}", f"_Generated {datetime.now():%Y-%m-%d %H:%M}. Local analysis only unless stated._", "",
         "## 1. Data used", f"- Samples: {counts['samples']} ({counts['samples_in_learning']} included in learning)",
         f"- Approved examples: {counts['approved_examples']}; rejected examples: {counts['rejected_examples']}",
         f"- Feedback decisions: {counts['feedback']}; drafts: {counts['drafts']}; revisions: {counts['revisions']}", ""]
    L += ["## 2. Observed features (measured from your samples)",
          "_Measurements of the text only. Small samples give unstable values. They do not prove authorship or ability._", ""]
    genres = sorted({s["genre"] for s in repo.list_samples(profile_id, learning_only=True)})
    if not genres:
        L.append("No learning-eligible samples yet.\n")
    for g in genres:
        gp = build_genre_profile(repo.list_samples(profile_id, genre=g, learning_only=True))
        f = gp["features"]
        L += [f"### {g}", f"- Samples: {gp['sample_count']}, words: {gp['word_count']}, evidence level: **{gp['evidence_level']}** (a label based on amount of text, not a probability)",
              f"- Mean sentence length: {f['mean_sentence_length']} words (SD {f['sentence_length_sd']}); mean paragraph length: {f['mean_paragraph_length']} words",
              f"- Vocabulary diversity (MATTR-50): {f['mattr_50']}; transitions per 100 words: {f['transitions_per_100_words']}; hedging per 100 words: {f['hedging_per_100_words']}",
              f"- Common transitions: {', '.join(f'{t} ({c})' for t, c in f['top_transitions'][:5]) or 'none found'}", ""]
    L += ["## 3. Preferences", "_Stated preferences come from you. Learned ones are estimates from your feedback and can be edited or reset._", ""]
    onboarding = repo.get_onboarding(profile_id)
    for k, v in onboarding.items():
        L.append(f"- [stated] {k}: {v}")
    for p in repo.get_preferences(profile_id):
        if p["value"]:
            L.append(f"- [{p['source']}] {p['genre']}: {DIMENSIONS.get(p['key'], {}).get('text', p['key'])} -> {p['value']} (evidence events: {p['evidence_n']})")
    avoid = [a["phrase"] for a in repo.list_avoid_phrases(profile_id)]
    if avoid:
        L.append("- [stated] Avoid: " + "; ".join(avoid))
    L.append(f"- Current preference-profile version: {repo.latest_preference_version(profile_id)}\n")
    fb = repo.list_feedback(profile_id)
    L += ["## 4. Feedback summary", f"- Decisions recorded: {len(fb)} (accept {sum(1 for e in fb if e['action']=='accept')}, edit {sum(1 for e in fb if e['action']=='edit')}, reject {sum(1 for e in fb if e['action']=='reject')})"]
    revs = repo.list_recent_revisions(profile_id, 200)
    rated = [r["style_rating"] for r in revs if r["style_rating"]]
    sat = [r["satisfaction_rating"] for r in revs if r["satisfaction_rating"]]
    if rated or sat:
        L.append(f"- User ratings (1-5): style similarity avg {sum(rated)/len(rated):.1f} (n={len(rated)}), satisfaction avg {sum(sat)/len(sat):.1f} (n={len(sat)})" if rated and sat else
                 f"- User ratings recorded: style n={len(rated)}, satisfaction n={len(sat)}")
    L.append("")
    L += ["## 5. Argument-analysis findings (saved reports)"]
    arg = repo.list_reports(profile_id, "argument", 3)
    L += [f"### {r['title']} ({r['created_at'][:16]})\n{r['content_md']}\n" for r in arg] or ["None saved yet.\n"]
    L += ["## 6. Recent revisions"]
    shown = 0
    for r in revs:
        if r["kind"] != "generated" or shown >= max_revisions:
            continue
        d = repo.get_draft(profile_id, r["draft_id"])
        if not d:
            continue
        expl = json.loads(r["explanation_json"] or "{}")
        cs = change_summary(d["original_text"], r["text"])
        L += [f"### {d['title']} ({d['genre']}, mode: {r['mode']}, status: {r['status']})",
              f"Words {cs['words_original']} -> {cs['words_revised']} ({cs['percent_change']}%). Provider: {r['provider']}.",
              "**Original**", "> " + d["original_text"].replace("\n", "\n> "), "**Revision (model-generated)**", "> " + r["text"].replace("\n", "\n> ")]
        for c in expl.get("changes", [])[:6]:
            L.append(f"- {c.get('change','')} ({c.get('reason','')})")
        L.append("")
        shown += 1
    if not shown:
        L.append("None yet.\n")
    L += ["## 7. Evaluation results (saved)"]
    ev = repo.list_reports(profile_id, "evaluation", 1)
    L += [ev[0]["content_md"]] if ev else ["None saved yet. Evaluation on synthetic cases is a pipeline check, not evidence of quality.\n"]
    return "\n".join(L)


def markdown_to_text(md: str) -> str:
    md = re.sub(r"^#+\s*", "", md, flags=re.M)
    return re.sub(r"[*_`]", "", md)


def markdown_to_docx_bytes(md: str) -> bytes:
    from docx import Document
    doc = Document()
    for line in md.splitlines():
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("#"):
            level = min(len(s) - len(s.lstrip("#")), 3)
            doc.add_heading(s.lstrip("# ").strip(), level=level)
        elif s.startswith("- "):
            doc.add_paragraph(re.sub(r"[*_`]", "", s[2:]), style="List Bullet")
        elif s.startswith(">"):
            doc.add_paragraph(re.sub(r"[*_`]", "", s.lstrip("> ")), style="Intense Quote")
        else:
            doc.add_paragraph(re.sub(r"[*_`]", "", s))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_preferences_json(repo: Repository, profile_id: int) -> str:
    return json.dumps({"onboarding": repo.get_onboarding(profile_id), "preferences": repo.get_preferences(profile_id),
                       "avoid_phrases": [a["phrase"] for a in repo.list_avoid_phrases(profile_id)],
                       "versions": [{"version": v["version"], "note": v["change_note"], "created_at": v["created_at"]} for v in repo.list_preference_versions(profile_id)]}, indent=2)
