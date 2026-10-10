"""Extract text from TXT, DOCX and PDF uploads with validation. Never returns empty text silently."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = (".txt", ".docx", ".pdf")


class ExtractionError(Exception):
    """User-presentable extraction failure."""


@dataclass
class Extraction:
    text: str
    warnings: list = field(default_factory=list)


def extract_text(filename: str, data: bytes, max_mb: int = 5) -> Extraction:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type '{ext or 'unknown'}'. Supported: TXT, DOCX, PDF. (Old .doc files must be saved as .docx first.)")
    if not data:
        raise ExtractionError("The file is empty.")
    if len(data) > max_mb * 1024 * 1024:
        raise ExtractionError(f"File is larger than the {max_mb} MB limit.")
    warnings: list[str] = []
    if ext == ".txt":
        if b"\x00" in data[:2000]:
            raise ExtractionError("This does not look like a plain-text file.")
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(enc)
                if enc != "utf-8-sig":
                    warnings.append("Decoded as Windows-1252 because the file was not UTF-8. Check special characters.")
                break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode("latin-1")
            warnings.append("Encoding was uncertain; decoded as Latin-1. Check special characters.")
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractionError("This file is not a valid .docx document.")
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionError("Could not read the .docx file. It may be corrupted or password protected.") from exc
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n\n".join(p for p in parts if p.strip())
        if doc.tables:
            warnings.append("Table contents were appended as text rows; check the order.")
    else:
        if not data.startswith(b"%PDF"):
            raise ExtractionError("This file is not a valid PDF.")
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionError("The PDF is encrypted. Remove the password first.")
            pages = [(p.extract_text() or "") for p in reader.pages]
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("Could not read the PDF. It may be corrupted.") from exc
        text = "\n\n".join(pages)
        empty = sum(1 for p in pages if not p.strip())
        if empty and empty < len(pages):
            warnings.append(f"{empty} page(s) had no extractable text (possibly images).")
        warnings.append("PDF extraction can break lines, headers and hyphenation. Review and correct the text below before saving.")
    text = text.replace("\x00", "").strip()
    if len(text) < 20:
        raise ExtractionError("No usable text could be extracted. If this is a scanned or image-only PDF, OCR is not supported; paste the text instead.")
    return Extraction(text, warnings)
